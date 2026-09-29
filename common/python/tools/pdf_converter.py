from __future__ import annotations

import os
from pathlib import Path
from typing import List, Dict, Any, Optional
from utils.logging import get_logger

LOGGER = get_logger("blinky.pdf_converter")


def convert_images_to_pdf(image_paths: List[str | Path], output_pdf: str | Path) -> Path:
    """Combines one or more images (PNG, JPG, WEBP) into a single PDF document."""
    from PIL import Image

    out_path = Path(output_pdf).resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)

    loaded_images = []
    for p in image_paths:
        img_path = Path(p).resolve()
        if not img_path.exists():
            continue
        img = Image.open(img_path)
        if img.mode != "RGB":
            img = img.convert("RGB")
        loaded_images.append(img)

    if not loaded_images:
        raise ValueError("No valid images provided for PDF conversion.")

    first_img = loaded_images[0]
    rest_imgs = loaded_images[1:]
    first_img.save(str(out_path), "PDF", resolution=100.0, save_all=True, append_images=rest_imgs)
    LOGGER.info(f"Successfully converted {len(loaded_images)} image(s) to PDF: {out_path}")
    return out_path


def convert_pdf_to_images(pdf_path: str | Path, output_dir: str | Path, fmt: str = "png") -> List[Path]:
    """Renders PDF pages as high-resolution image files (PNG/JPG)."""
    pdf_file = Path(pdf_path).resolve()
    out_dir = Path(output_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    saved_images: List[Path] = []

    # Try PyMuPDF (fitz) first - ultra fast
    try:
        import fitz
        doc = fitz.open(str(pdf_file))
        for idx, page in enumerate(doc):
            pix = page.get_pixmap(dpi=150)
            img_name = f"{pdf_file.stem}_page_{idx + 1}.{fmt}"
            img_path = out_dir / img_name
            pix.save(str(img_path))
            saved_images.append(img_path)
        LOGGER.info(f"Rendered {len(saved_images)} PDF page(s) to {fmt.upper()} using PyMuPDF.")
        return saved_images
    except ImportError:
        pass

    # Fallback to pdf2image
    try:
        from pdf2image import convert_from_path
        images = convert_from_path(str(pdf_file), dpi=150)
        for idx, img in enumerate(images):
            img_name = f"{pdf_file.stem}_page_{idx + 1}.{fmt}"
            img_path = out_dir / img_name
            img.save(str(img_path), fmt.upper())
            saved_images.append(img_path)
        LOGGER.info(f"Rendered {len(saved_images)} PDF page(s) using pdf2image.")
        return saved_images
    except Exception as exc:
        raise RuntimeError(f"Failed rendering PDF to images: {exc}")


def convert_txt_to_pdf(txt_path: str | Path, output_pdf: str | Path) -> Path:
    """Converts a TXT or Markdown file to a formatted PDF document."""
    txt_file = Path(txt_path).resolve()
    out_pdf = Path(output_pdf).resolve()
    out_pdf.parent.mkdir(parents=True, exist_ok=True)

    content = txt_file.read_text(encoding="utf-8", errors="ignore")

    # Try ReportLab first
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

        doc = SimpleDocTemplate(str(out_pdf), pagesize=letter)
        styles = getSampleStyleSheet()
        normal_style = styles["Normal"]
        normal_style.fontSize = 10
        normal_style.leading = 14

        story = []
        for line in content.splitlines():
            clean_line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if not clean_line.strip():
                story.append(Spacer(1, 8))
            else:
                story.append(Paragraph(clean_line, normal_style))

        doc.build(story)
        LOGGER.info(f"Converted TXT to PDF using ReportLab: {out_pdf}")
        return out_pdf
    except ImportError:
        pass

    # Fallback to PIL (Pillow) image to PDF
    try:
        from PIL import Image, ImageDraw
        lines = content.splitlines() or [""]
        line_height = 22
        padding = 40
        width = 612  # Standard Letter width in points
        height = max(792, padding * 2 + len(lines) * line_height)
        img = Image.new("RGB", (width, height), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)

        y = padding
        for line in lines:
            draw.text((padding, y), line, fill=(0, 0, 0))
            y += line_height

        img.save(str(out_pdf), "PDF", resolution=100.0)
        LOGGER.info(f"Converted TXT to PDF using PIL fallback: {out_pdf}")
        return out_pdf
    except Exception as exc:
        raise RuntimeError(f"Failed converting TXT to PDF: {exc}")


def convert_docx_to_pdf(docx_path: str | Path, output_pdf: str | Path) -> Path:
    """Converts a Word DOCX file to PDF."""
    docx_file = Path(docx_path).resolve()
    out_pdf = Path(output_pdf).resolve()
    out_pdf.parent.mkdir(parents=True, exist_ok=True)

    # 1. Try docx2pdf (Native MS Word / COM interface on Windows)
    try:
        from docx2pdf import convert
        convert(str(docx_file), str(out_pdf))
        LOGGER.info(f"Converted DOCX to PDF using docx2pdf: {out_pdf}")
        return out_pdf
    except Exception:
        pass

    # 2. Fallback: Parse DOCX text via python-docx and output formatted PDF
    try:
        import docx
        doc_obj = docx.Document(str(docx_file))
        paragraphs = [p.text for p in doc_obj.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)

        temp_txt = docx_file.with_suffix(".temp.txt")
        temp_txt.write_text(full_text, encoding="utf-8")
        result = convert_txt_to_pdf(temp_txt, out_pdf)
        try:
            temp_txt.unlink()
        except Exception:
            pass
        LOGGER.info(f"Converted DOCX text to PDF via python-docx fallback: {out_pdf}")
        return result
    except Exception as exc:
        raise RuntimeError(f"Failed converting DOCX to PDF: {exc}")


def convert_pdf_to_docx(pdf_path: str | Path, output_docx: str | Path) -> Path:
    """Converts a PDF file to editable Word DOCX format."""
    pdf_file = Path(pdf_path).resolve()
    out_docx = Path(output_docx).resolve()
    out_docx.parent.mkdir(parents=True, exist_ok=True)

    # 1. Try pdf2docx
    try:
        from pdf2docx import Converter
        cv = Converter(str(pdf_file))
        cv.convert(str(out_docx), start=0, end=None)
        cv.close()
        LOGGER.info(f"Converted PDF to DOCX using pdf2docx: {out_docx}")
        return out_docx
    except ImportError:
        pass

    # 2. Fallback: Extract text and generate python-docx file
    try:
        import docx
        import fitz
        pdf_doc = fitz.open(str(pdf_file))
        doc_obj = docx.Document()

        for page in pdf_doc:
            text = page.get_text()
            if text.strip():
                doc_obj.add_paragraph(text.strip())

        doc_obj.save(str(out_docx))
        LOGGER.info(f"Converted PDF to DOCX using fitz + python-docx fallback: {out_docx}")
        return out_docx
    except Exception as exc:
        raise RuntimeError(f"Failed converting PDF to DOCX: {exc}")


def convert_document(
    input_path: str | Path | List[str | Path],
    target_format: str,
    output_path: Optional[str | Path] = None
) -> Dict[str, Any]:
    """
    Unified format conversion dispatcher.
    Supports: pdf, docx, txt, png, jpg, images.
    """
    target = target_format.lower().lstrip(".")

    if isinstance(input_path, list):
        # Images to PDF batch conversion
        out_file = Path(output_path or "combined_output.pdf").resolve()
        result_path = convert_images_to_pdf(input_path, out_file)
        return {"success": True, "output_path": str(result_path), "format": "pdf"}

    in_file = Path(input_path).resolve()
    if not in_file.exists():
        raise FileNotFoundError(f"Input file not found: {in_file}")

    source_ext = in_file.suffix.lstrip(".").lower()

    if target == "pdf":
        if source_ext in {"png", "jpg", "jpeg", "webp"}:
            out_file = Path(output_path or in_file.with_suffix(".pdf")).resolve()
            res = convert_images_to_pdf([in_file], out_file)
        elif source_ext == "docx":
            out_file = Path(output_path or in_file.with_suffix(".pdf")).resolve()
            res = convert_docx_to_pdf(in_file, out_file)
        else:
            out_file = Path(output_path or in_file.with_suffix(".pdf")).resolve()
            res = convert_txt_to_pdf(in_file, out_file)
        return {"success": True, "output_path": str(res), "format": "pdf"}

    elif target in {"docx", "doc"}:
        out_file = Path(output_path or in_file.with_suffix(".docx")).resolve()
        res = convert_pdf_to_docx(in_file, out_file)
        return {"success": True, "output_path": str(res), "format": "docx"}

    elif target in {"png", "jpg", "jpeg"}:
        out_dir = Path(output_path or in_file.parent / f"{in_file.stem}_pages").resolve()
        res_list = convert_pdf_to_images(in_file, out_dir, fmt=target)
        return {"success": True, "output_paths": [str(p) for p in res_list], "format": target}

    else:
        raise ValueError(f"Unsupported target format conversion: {source_ext} -> {target}")
